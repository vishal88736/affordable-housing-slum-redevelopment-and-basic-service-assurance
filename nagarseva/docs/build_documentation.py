"""Build the maintained NagarSeva Markdown documentation as a Word document."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "NagarSeva_Project_Documentation.md"
OUTPUT = ROOT / "docs" / "NagarSeva_Project_Documentation.docx"


def add_table(document: Document, lines: list[str]) -> None:
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(set(cell) <= {"-", ":", " "} for cell in row)]
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Light Shading Accent 1"
    for index, value in enumerate(rows[0]):
        table.rows[0].cells[index].text = value
    for row in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def add_code(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="No Spacing")
    paragraph.paragraph_format.left_indent = Inches(.3)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(35, 70, 85)


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(.65)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.75)
    section.right_margin = Inches(.75)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(35, 60, 72)
    for name, size, colour in [("Title", 26, RGBColor(9, 42, 67)), ("Heading 1", 18, RGBColor(9, 42, 67)), ("Heading 2", 13, RGBColor(21, 130, 130)), ("Heading 3", 11, RGBColor(23, 60, 84))]:
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = colour

    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(line[2:])
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("#### "):
            document.add_heading(line[5:], level=3)
        elif line.startswith("```"):
            code_lines = []
            index += 1
            while index < len(lines) and lines[index].strip() != "```":
                code_lines.append(lines[index])
                index += 1
            add_code(document, "\n".join(code_lines))
        elif line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue
        elif line.startswith("> "):
            paragraph = document.add_paragraph(style="Intense Quote")
            paragraph.add_run(line[2:])
        elif re.match(r"^\d+\. ", line):
            document.add_paragraph(re.sub(r"^\d+\. ", "", line), style="List Number")
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            paragraph = document.add_paragraph()
            paragraph.add_run(line)
        index += 1

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("NagarSeva · Maintained project documentation · Synthetic demo data").font.size = Pt(8)
    document.core_properties.title = "NagarSeva Project Documentation"
    document.core_properties.subject = "Architecture, implementation and maintenance guide"
    document.core_properties.author = "NagarSeva engineering team"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
